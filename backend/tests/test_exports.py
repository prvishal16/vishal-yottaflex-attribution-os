from __future__ import annotations

import csv
import io
import json
import zipfile
from decimal import Decimal

import duckdb
from openpyxl import load_workbook


def _csv_rows(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.text)))


def _json(response) -> dict:
    return json.loads(response.text)


def test_overview_export_covers_every_month(client) -> None:
    response = client.get("/api/exports/overview?format=csv")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    assert "executive-overview" in response.headers["content-disposition"]
    rows = _csv_rows(response)
    assert rows[0] == [
        "month",
        "attributed_revenue",
        "eligible_revenue",
        "spend",
        "attributed_conversions",
    ]
    assert len(rows) == 14
    assert rows[0][0] != rows[1][0]

    structured = client.get("/api/exports/overview?format=json")
    payload = _json(structured)
    assert payload["context"]["report"] == "Executive overview"
    assert payload["context"]["artifact_type"] == "analytical_export"
    assert payload["context"]["contains_ground_truth"] is False
    assert len(payload["rows"]) == 13
    names = [sheet["name"] for sheet in payload["sheets"]]
    assert {"Monthly", "Channel contribution", "Model comparison", "Top campaigns"} <= set(names)

    workbook = load_workbook(io.BytesIO(client.get("/api/exports/overview?format=xlsx").content), read_only=True)
    assert workbook.sheetnames[0] == "Summary"
    assert "Monthly" in workbook.sheetnames
    workbook.close()

    pdf = client.get("/api/exports/overview?format=pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")


def test_campaign_export_is_scoped_and_consistent_across_formats(client) -> None:
    campaign_id = client.get("/api/campaigns?limit=1").json()["items"][0]["campaign_id"]
    detail = client.get(f"/api/campaigns/{campaign_id}?model=linear").json()
    response = client.get(f"/api/exports/campaign/{campaign_id}?model=linear&format=json")
    assert response.status_code == 200
    payload = _json(response)
    assert payload["context"]["campaign_id"] == campaign_id
    assert payload["context"]["model"] == "linear"
    assert len(payload["rows"]) == len(detail["monthly"])
    assert payload["campaign"]["campaign"]["campaign_id"] == campaign_id
    assert all(row["month"] for row in payload["rows"])

    csv_rows = _csv_rows(
        client.get(f"/api/exports/campaign/{campaign_id}?model=linear&format=csv")
    )
    assert len(csv_rows) == len(detail["monthly"]) + 1

    workbook = load_workbook(
        io.BytesIO(
            client.get(f"/api/exports/campaign/{campaign_id}?model=linear&format=xlsx").content
        ),
        read_only=True,
    )
    assert {"Summary", "Monthly credit", "Model comparison", "Cost records"} <= set(
        workbook.sheetnames
    )
    monthly = workbook["Monthly credit"]
    values = list(monthly.iter_rows(min_row=2, values_only=True))
    assert values
    assert all(isinstance(row[1], (int, float)) for row in values)
    workbook.close()

    pdf = client.get(f"/api/exports/campaign/{campaign_id}?model=linear&format=pdf")
    assert pdf.content.startswith(b"%PDF-")
    assert client.get("/api/exports/campaign/CMP-2026-999999?format=csv").status_code == 404


def test_cost_export_only_contains_the_campaign_records(client) -> None:
    campaign_id = client.get("/api/campaigns?limit=1").json()["items"][0]["campaign_id"]
    detail = client.get(f"/api/campaigns/{campaign_id}?model=linear").json()
    response = client.get(f"/api/exports/cost-records?campaign_id={campaign_id}&format=csv")
    assert response.status_code == 200
    rows = _csv_rows(response)
    assert rows[0][0] == "cost_id"
    assert len(rows) == len(detail["cost_records"]) + 1
    expected_ids = {record["cost_id"] for record in detail["cost_records"]}
    assert {row[0] for row in rows[1:]} == expected_ids

    amounts = [Decimal(row[2]) for row in rows[1:]]
    detail_total = sum(
        Decimal(str(record["amount"])) for record in detail["cost_records"]
    )
    assert sum(amounts) == detail_total
    pdf = client.get(f"/api/exports/cost-records?campaign_id={campaign_id}&format=pdf")
    assert pdf.content.startswith(b"%PDF-")
    assert client.get("/api/exports/cost-records?campaign_id=CMP-2026-999999").status_code == 404


def test_exception_export_returns_all_matching_rows_not_the_page(client) -> None:
    queue = client.get("/api/exceptions?limit=5").json()
    assert queue["total"] == 10
    response = client.get("/api/exports/exceptions?format=json")
    payload = _json(response)
    assert payload["context"]["artifact_type"] == "detection_output"
    assert payload["context"]["contains_ground_truth"] is False
    assert payload["context"]["row_count"] == 10
    assert len(payload["rows"]) == 10
    assert len(payload["rows"]) > queue["limit"]

    high = client.get("/api/exports/exceptions?format=json&severity=HIGH").json()
    assert high["context"]["row_count"] == 6
    assert {row["severity"] for row in high["rows"]} == {"HIGH"}

    typed = client.get(
        "/api/exports/exceptions?format=json&issue_type=DUPLICATE_REVENUE"
    ).json()
    assert typed["context"]["row_count"] == 2
    assert {row["issue_type"] for row in typed["rows"]} == {"DUPLICATE_REVENUE"}

    searched = client.get("/api/exports/exceptions?format=json&search=source%20reference").json()
    assert searched["context"]["row_count"] > 0
    assert all("source reference" in row["explanation"].lower() for row in searched["rows"])

    empty = client.get("/api/exports/exceptions?format=csv&search=no-such-issue-xyz")
    assert empty.status_code == 200
    assert len(_csv_rows(empty)) == 1
    empty_pdf = client.get("/api/exports/exceptions?format=pdf&search=no-such-issue-xyz")
    assert empty_pdf.content.startswith(b"%PDF-")


def test_validation_export_filters_without_losing_the_summary(client) -> None:
    report = client.get("/api/validation/report").json()
    warnings = [check for check in report["checks"] if check["status"] == "WARNING"]
    response = client.get("/api/exports/validation?format=json&status=WARNING")
    payload = _json(response)
    assert payload["context"]["row_count"] == len(warnings)
    assert len(payload["rows"]) == len(warnings)
    assert {row["status"] for row in payload["rows"]} == {"WARNING"}
    categories = next(sheet for sheet in payload["sheets"] if sheet["name"] == "Categories")
    assert categories["row_count"] >= 5

    filtered = client.get("/api/exports/validation?format=csv&category=FINANCIAL")
    rows = _csv_rows(filtered)
    assert len(rows) == 9
    assert all(row[1] == "FINANCIAL" for row in rows[1:])

    searched = client.get("/api/exports/validation?format=json&search=duplicate")
    assert searched.json()["context"]["row_count"] > 0


def test_evaluation_export_is_the_only_explicit_ground_truth_artifact(client) -> None:
    response = client.get("/api/exports/evaluation?format=json")
    assert response.status_code == 200
    payload = _json(response)
    assert payload["context"]["artifact_type"] == "evaluation_artifact"
    assert payload["context"]["contains_ground_truth"] is True
    assert "ground-truth" in payload["context"]["notice"].lower()
    assert payload["evaluation"]["totals"]["ground_truth"] == 10
    assert len(payload["evaluation"]["matches"]) == 10
    assert any(sheet["name"] == "Ground truth matches" for sheet in payload["sheets"])

    workbook = load_workbook(
        io.BytesIO(client.get("/api/exports/evaluation?format=xlsx").content), read_only=True
    )
    assert "Ground truth matches" in workbook.sheetnames
    workbook.close()

    pdf = client.get("/api/exports/evaluation?format=pdf")
    assert pdf.content.startswith(b"%PDF-")

    by_type_csv = _csv_rows(client.get("/api/exports/evaluation?format=csv"))
    assert by_type_csv[0][0] == "issue_type"
    assert by_type_csv[0][1] == "ground_truth"
    assert len(by_type_csv) == 6


def test_evidence_export_covers_every_detected_issue(client) -> None:
    issues = client.get("/api/exceptions?limit=100").json()["items"]
    assert len(issues) == 10
    for issue in issues:
        response = client.get(f"/api/exports/evidence/{issue['issue_id']}?format=json")
        assert response.status_code == 200, issue["issue_id"]
        payload = _json(response)
        assert payload["packet"]["issue_id"] == issue["issue_id"]
        assert payload["packet"]["related"]["disposition_state"]["applied"] is False
        assert payload["context"]["contains_ground_truth"] is False
    assert client.get("/api/exports/evidence/ISS-2026-999999").status_code == 404


def test_journey_export_contains_every_candidate_touchpoint(client) -> None:
    journey = client.get("/api/demo/journey").json()
    reference = journey["reference"]
    response = client.get(f"/api/exports/journey/{reference}?format=json")
    assert response.status_code == 200
    payload = _json(response)
    assert payload["context"]["row_count"] == len(journey["touchpoints"])
    assert len(payload["rows"]) == len(journey["touchpoints"])
    allocations = next(sheet for sheet in payload["sheets"] if sheet["name"] == "Allocations")
    expected_allocations = sum(
        len(touch.get("credits", {})) for touch in journey["touchpoints"]
    )
    assert allocations["row_count"] == expected_allocations
    models = next(sheet for sheet in payload["sheets"] if sheet["name"] == "Models")
    assert {row["model"] for row in models["rows"]} == {
        "first_touch",
        "last_touch",
        "linear",
    }
    assert all(row["persisted_match"] for row in models["rows"])
    assert client.get("/api/exports/journey/OPP-2026-999999").status_code == 404

    csv_rows = _csv_rows(client.get(f"/api/exports/journey/{reference}?format=csv"))
    assert len(csv_rows) == len(journey["touchpoints"]) + 1


def test_journey_export_handles_no_recognized_conversion(client, run_paths: dict) -> None:
    connection = duckdb.connect(str(run_paths["db"]), read_only=True)
    try:
        order_id = connection.execute(
            "SELECT order_id FROM orders "
            "WHERE status NOT IN ('PAID', 'PARTIALLY_REFUNDED') ORDER BY order_id LIMIT 1"
        ).fetchone()[0]
    finally:
        connection.close()

    journey = client.get(f"/api/journeys/{order_id}").json()
    assert journey["status"] == "NO_RECOGNIZED_CONVERSION"
    assert journey.get("conversion") is None

    response = client.get(f"/api/exports/journey/{order_id}?format=json")
    assert response.status_code == 200
    payload = _json(response)
    assert payload["context"]["row_count"] == 0
    assert payload["journey"]["status"] == "NO_RECOGNIZED_CONVERSION"
    conversion = next(sheet for sheet in payload["sheets"] if sheet["name"] == "Conversion")
    assert {row["field"]: row["value"] for row in conversion["rows"]}["message"]
    assert client.get(f"/api/exports/journey/{order_id}?format=csv").status_code == 200
    assert client.get(f"/api/exports/journey/{order_id}?format=pdf").content.startswith(b"%PDF-")


def test_run_bundle_and_sql_package_exclude_ground_truth(client, bundle: dict) -> None:
    run_id = bundle["config"].run_id
    response = client.get(f"/api/exports/run/{run_id}?format=zip")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = set(archive.namelist())
        assert {
            "README.txt",
            "manifest.json",
            "run_config.json",
            "evaluation_summary.json",
            "counts.csv",
            "validation_checks.csv",
            "validation_summary.json",
        } <= names
        manifest = archive.read("manifest.json").decode("utf-8")
        evaluation = archive.read("evaluation_summary.json").decode("utf-8")
        counts = list(csv.reader(io.StringIO(archive.read("counts.csv").decode("utf-8"))))
        assert "ground_truth" not in manifest
        assert "expected_record_ids" not in evaluation
        assert '"matches"' not in evaluation
        assert any(row[0] == "campaigns" for row in counts[1:])
        checks = archive.read("validation_checks.csv").decode("utf-8")
        assert len(checks.splitlines()) == 120

    sql = client.get(f"/api/exports/run/{run_id}?format=sql")
    assert sql.status_code == 200
    assert sql.content[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(sql.content)) as archive:
        names = set(archive.namelist())
        assert {"schema.sql", "load.sql", "README.txt"} <= names
        load = archive.read("load.sql").decode("utf-8")
        assert "FROM '" in load
        assert "AppData" not in load
        assert "C:/" not in load
        assert "CREATE TABLE" in archive.read("schema.sql").decode("utf-8")

    json_export = client.get(f"/api/exports/run/{run_id}?format=json")
    assert json_export.status_code == 200
    payload = _json(json_export)
    assert payload["context"]["artifact_type"] == "run_manifest"
    assert payload["context"]["contains_ground_truth"] is False
    assert payload["context"]["combined_hash"] == bundle["manifest"]["combined_hash"]


def test_export_catalog_and_error_handling(client, bundle: dict) -> None:
    body = client.get("/api/exports/formats").json()
    assert [item["format"] for item in body["formats"]] == ["csv", "json", "xlsx", "pdf"]
    assert [item["format"] for item in body["run_formats"]] == [
        "json",
        "csv",
        "xlsx",
        "pdf",
        "sql",
        "zip",
    ]
    assert body["run_bundle"]["contains_ground_truth"] is False

    assert client.get("/api/exports/overview?format=docx").status_code == 422
    assert client.get("/api/exports/run/missing-123?format=json").status_code == 404
    assert client.get(f"/api/exports/run/{bundle['config'].run_id}?format=docx").status_code == 422
